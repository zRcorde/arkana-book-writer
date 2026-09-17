"""
Arkana Book Writer — local ebook studio (PDF/EPUB).
https://github.com/ (Arkana Book Writer project)
Support: buymeacoffee.com/re_code | USDC (Solana): 4qroECsHYTqk7Sda412LiiXofuMj5xhYRHmrKqXbKM8X | BTC: bc1puezkfpq7eea6whzrh68qmhvw9gmxynwpe8mr5w509ycqcn3y64hq9ywjrn
Partnerships / contact: contact@rewebfolio.xyz
"""
import os
import markdown
from html import escape as html_escape
from docx import Document
from pypdf import PdfReader
from bs4 import BeautifulSoup
from .models import EbookContent

# Sentinel placeholders for text the parser can't extract (no author metadata,
# no heading found). Resolved to a localized string in web_app.py after
# parsing, once the target export language is known.
UNKNOWN_AUTHOR = "__UNKNOWN_AUTHOR__"
INTRO_TITLE = "__INTRO__"
MAIN_CONTENT_TITLE = "__MAIN_CONTENT__"
EXTRACTED_CONTENT_TITLE = "__EXTRACTED_CONTENT__"

class BaseParser:
    def parse(self, file_path: str) -> EbookContent:
        raise NotImplementedError

class MarkdownParser(BaseParser):
    def parse(self, file_path: str) -> EbookContent:
        with open(file_path, 'r', encoding='utf-8') as f:
            text = f.read()
        
        html = markdown.markdown(text, extensions=['extra', 'toc'])
        soup = BeautifulSoup(html, 'html.parser')
        
        content = EbookContent(title=os.path.basename(file_path).split('.')[0], author=UNKNOWN_AUTHOR)
        
        # Simple splitting by H1 or H2 for chapters
        headers = soup.find_all(['h1', 'h2'])
        if not headers:
            content.add_chapter(MAIN_CONTENT_TITLE, html)
        else:
            current_chapter_title = INTRO_TITLE
            current_chapter_content = []
            
            for element in soup.contents:
                if element.name in ['h1', 'h2']:
                    if current_chapter_content:
                        content.add_chapter(current_chapter_title, "".join(str(e) for e in current_chapter_content))
                    current_chapter_title = html_escape(element.get_text())
                    current_chapter_content = []
                else:
                    current_chapter_content.append(element)
            
            if current_chapter_content:
                content.add_chapter(current_chapter_title, "".join(str(e) for e in current_chapter_content))
                
        return content

class DocxParser(BaseParser):
    def parse(self, file_path: str) -> EbookContent:
        doc = Document(file_path)
        content = EbookContent(title=os.path.basename(file_path).split('.')[0], author=UNKNOWN_AUTHOR)
        
        current_chapter_title = INTRO_TITLE
        current_chapter_html = []
        
        for para in doc.paragraphs:
            if para.style.name.startswith('Heading'):
                if current_chapter_html:
                    content.add_chapter(current_chapter_title, "".join(current_chapter_html))
                current_chapter_title = html_escape(para.text)
                current_chapter_html = []
            else:
                if para.text.strip():
                    # Escape so raw <, >, & in the source text can't break
                    # the HTML the editor/preview/exporter renders it into.
                    current_chapter_html.append(f"<p>{html_escape(para.text)}</p>")
        
        if current_chapter_html:
            content.add_chapter(current_chapter_title, "".join(current_chapter_html))
            
        return content

class PdfParser(BaseParser):
    def parse(self, file_path: str) -> EbookContent:
        reader = PdfReader(file_path)
        content = EbookContent(title=os.path.basename(file_path).split('.')[0], author=UNKNOWN_AUTHOR)
        
        full_text = ""
        for page in reader.pages:
            full_text += page.extract_text() + "\n\n"

        # Escape before turning newlines into <br>, so raw <, >, & extracted
        # from the PDF text can't break the HTML it gets embedded into.
        # PDF is harder to split into chapters, so we treat it as one for now or split by double newlines
        content.add_chapter(EXTRACTED_CONTENT_TITLE, html_escape(full_text).replace('\n', '<br>'))
        
        return content

def get_parser(file_extension: str) -> BaseParser:
    parsers = {
        '.md': MarkdownParser(),
        '.docx': DocxParser(),
        '.pdf': PdfParser()
    }
    return parsers.get(file_extension.lower())
